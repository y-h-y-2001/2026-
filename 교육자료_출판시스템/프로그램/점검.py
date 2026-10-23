#!/usr/bin/env python3
"""샘플 자료로 추출, 중복 검사, 원고 만들기가 되는지 확인한다."""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from edupub.build_db import build
from edupub.db import add_item, connect
from edupub.dedupe import find_duplicates
from edupub.docx_out import build_manuscripts
from edupub.ocr import ocr_status
from edupub.paths import ROOT, SAMPLE_DIR, VERIFY_DIR
from edupub.pdf_text import extract_pdf

EXPECTED_STATEMENTS = {
    "9역01-01": "역사와 역사 탐구의 의미를 파악하고, 역사 학습의 목적을 다각도로 탐색한다.",
    "9사(지리)01-01": "세계 여러 지역의 특성을 해당 지역의 위치와 자연⋅인문환경을 고려하여 추론한다.",
    "10한사1-01-01": "고대 국가의 형성과 성장 과정을 파악한다.",
    "4사01-01": "주변 여러 장소에서의 경험과 느낌을 다양한 방식으로 표현하고, 장소감을 나누며 서로 존중하는 태도를 지닌다.",
}

HISTORY_PUBLISHERS = {
    "(주)지학사",
    "(주)미래엔",
    "주식회사 리베르스쿨",
    "(주)비상교육",
    "(주)해냄에듀",
    "(주)천재교과서",
    "동아출판(주)",
}

FONT_CANDIDATES = [
    "/System/Library/Fonts/AppleSDGothicNeo.ttc",
    "/Library/Fonts/Arial Unicode.ttf",
    "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
    "C:/Windows/Fonts/malgun.ttf",
]


def main() -> None:
    lines: list[str] = []
    build()
    lines.append("데이터베이스를 원문에서 다시 만들었습니다.")
    connection = connect()
    try:
        for code, statement in EXPECTED_STATEMENTS.items():
            row = connection.execute("SELECT statement FROM standards WHERE code = ?", (code,)).fetchone()
            if row is None or row["statement"] != statement:
                raise SystemExit(f"성취기준 문장이 원문과 다릅니다: {code}")
        lines.append(f"성취기준 문장 {len(EXPECTED_STATEMENTS)}개를 고시 원문과 대조했습니다.")

        standard_count = connection.execute("SELECT COUNT(*) FROM standards").fetchone()[0]
        if standard_count < 300:
            raise SystemExit(f"성취기준 수가 너무 적습니다: {standard_count}")
        lines.append(f"성취기준 {standard_count}개를 넣었습니다.")

        history = {
            row["publisher"]
            for row in connection.execute("SELECT publisher FROM textbooks WHERE title = '역사 ①'")
        }
        if history != HISTORY_PUBLISHERS:
            raise SystemExit(f"중학교 역사 ① 발행사가 공고와 다릅니다: {sorted(history)}")
        author = connection.execute(
            "SELECT lead_author, publisher FROM textbooks WHERE approval_number = '중10525-P-01'"
        ).fetchone()
        if author["publisher"] != "동아출판(주)" or author["lead_author"] != "설규주":
            raise SystemExit(f"사회 ① 대표 저자 칸이 공고와 다릅니다: {tuple(author)}")
        lines.append("중학교 역사 ① 7종의 발행사와 사회 ① 동아출판 대표 저자 칸을 공고와 대조했습니다.")

        korean = connection.execute("SELECT COUNT(*) FROM textbooks WHERE title = '한국사1'").fetchone()[0]
        if korean != 9:
            raise SystemExit(f"한국사1 종수가 9가 아닙니다: {korean}")
        lines.append("고등학교 한국사1 9종을 확인했습니다.")

        secured = connection.execute("SELECT COUNT(*) FROM textbooks WHERE content_secured = 1").fetchone()[0]
        books = connection.execute("SELECT COUNT(*) FROM textbooks").fetchone()[0]
        if secured != 0:
            raise SystemExit("교과서 본문이 확보된 것처럼 표시된 도서가 있습니다.")
        lines.append(f"교과서 서지 {books}종은 모두 본문 미확보 상태입니다.")

        publishers = connection.execute("SELECT COUNT(DISTINCT publisher) FROM textbooks").fetchone()[0]
        if publishers < 8:
            raise SystemExit(f"발행사가 8곳 미만입니다: {publishers}")
        lines.append(f"확인된 발행사는 {publishers}곳입니다.")

        orphan = connection.execute(
            """
            SELECT COUNT(*) FROM items i
            LEFT JOIN standards s ON s.code = i.standard_code
            LEFT JOIN sources src ON src.source_id = i.source_id
            WHERE s.code IS NULL OR src.source_id IS NULL
            """
        ).fetchone()[0]
        if orphan:
            raise SystemExit("출처나 성취기준이 없는 문항이 있습니다.")
        lines.append("샘플 문항은 모두 성취기준과 출처에 연결되어 있습니다.")

        try:
            add_item(
                connection,
                item_id="SAMPLE-거절",
                stem="없는 성취기준을 가리키는 문항",
                choices=[],
                answer="없음",
                explanation="이건 저장되면 안 됩니다.",
                item_type="서술형",
                standard_code="없는코드-01",
                source_id="SRC-자체작성-샘플",
                origin="자체작성_샘플",
            )
            raise SystemExit("없는 성취기준 문항이 저장되었습니다.")
        except ValueError:
            lines.append("없는 성취기준을 가리키는 문항은 저장되지 않았습니다.")

        pdf_path = _sample_pdf()
        extracted = extract_pdf(pdf_path)
        extracted_text = extracted.read_text(encoding="utf-8")
        if "역사 탐구란 과거의 기록을 외우는 일만을 뜻한다." not in extracted_text:
            raise SystemExit("PDF에서 뽑은 글자가 원문과 다릅니다.")
        lines.append(f"PDF 글자 추출을 확인했습니다: {extracted.name}")

        pairs = find_duplicates(connection)
        exact = [pair for pair in pairs if pair.kind == "동일" and {pair.left_id, pair.right_id} == {"SAMPLE-001", "SAMPLE-002"}]
        similar = [pair for pair in pairs if pair.kind == "유사" and {pair.left_id, pair.right_id} == {"SAMPLE-001", "SAMPLE-003"}]
        if not exact or not similar:
            raise SystemExit(f"중복 검사 결과가 예상과 다릅니다: {pairs}")
        lines.append("같은 문항과 비슷한 문항을 구분했습니다.")

        student, teacher = build_manuscripts(connection, ["SAMPLE-001", "SAMPLE-004"])
        from docx import Document

        student_text = "\n".join(paragraph.text for paragraph in Document(student).paragraphs)
        teacher_text = "\n".join(paragraph.text for paragraph in Document(teacher).paragraphs)
        if "정답:" in student_text:
            raise SystemExit("학생용 원고에 정답이 들어 있습니다.")
        if "9역01-01" not in teacher_text or "교육부 고시 제2022-33호" not in teacher_text:
            raise SystemExit("교사용 원고에 성취기준 출처가 없습니다.")
        if "②" not in teacher_text:
            raise SystemExit("교사용 원고에 정답이 없습니다.")
        lines.append(f"학생용 원고: {student.relative_to(ROOT)}")
        lines.append(f"교사용 원고: {teacher.relative_to(ROOT)}")

        lines.append(f"OCR 상태: {ocr_status()}")
    finally:
        connection.close()

    VERIFY_DIR.mkdir(parents=True, exist_ok=True)
    report = VERIFY_DIR / "검증결과.md"
    report.write_text("# 검증 결과\n\n" + "\n".join(f"- {line}" for line in lines) + "\n", encoding="utf-8")
    print("\n".join(lines))
    print(report)


def _sample_pdf() -> Path:
    import pymupdf

    font = next((path for path in FONT_CANDIDATES if Path(path).exists()), "")
    if not font:
        raise SystemExit("한글을 표시할 글꼴을 찾지 못했습니다.")
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    path = SAMPLE_DIR / "샘플_점검.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_font(fontname="kr", fontfile=font)
    page.insert_text((50, 80), "역사 탐구란 과거의 기록을 외우는 일만을 뜻한다.", fontname="kr", fontsize=14)
    page.insert_text((50, 120), "이 문장은 점검용으로 직접 만든 문장이다.", fontname="kr", fontsize=12)
    document.save(path)
    document.close()
    return path


if __name__ == "__main__":
    main()
