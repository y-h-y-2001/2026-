"""학생용 원고와 교사용 원고를 워드 파일로 만든다."""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from edupub.db import now_seoul, require_source, require_standard
from edupub.paths import STUDENT_OUT, TEACHER_OUT


def _set_korean_font(run, size: int = 11, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")


def _add(document: Document, text: str, size: int = 11, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    _set_korean_font(run, size, bold)


def build_manuscripts(connection: sqlite3.Connection, item_ids: list[str] | None = None) -> tuple[Path, Path]:
    query = "SELECT * FROM items"
    params: tuple = ()
    if item_ids:
        marks = ",".join("?" for _ in item_ids)
        query += f" WHERE item_id IN ({marks})"
        params = tuple(item_ids)
    query += " ORDER BY item_id"
    items = connection.execute(query, params).fetchall()
    if not items:
        raise ValueError("원고로 만들 문항이 없습니다.")
    for item in items:
        require_standard(connection, item["standard_code"])
        require_source(connection, item["source_id"])

    stamp = now_seoul().replace(":", "").replace(" ", "_")
    STUDENT_OUT.mkdir(parents=True, exist_ok=True)
    TEACHER_OUT.mkdir(parents=True, exist_ok=True)
    student_path = STUDENT_OUT / f"학생용_{stamp}.docx"
    teacher_path = TEACHER_OUT / f"교사용_{stamp}.docx"
    _student(student_path, items, stamp)
    _teacher(connection, teacher_path, items, stamp)
    return student_path, teacher_path


def _student(path: Path, items: list, stamp: str) -> None:
    document = Document()
    _add(document, "학생용 학습 자료", 18, True)
    _add(document, f"문서 번호: {stamp}")
    _add(document, "문항의 정답과 출처는 교사용 원고에 있습니다.")
    for number, item in enumerate(items, start=1):
        _add(document, f"{number}. {item['stem']}", 12, True)
        for choice in json.loads(item["choices_json"]):
            _add(document, choice)
        if not json.loads(item["choices_json"]):
            _add(document, "답: ______________________________")
    document.save(path)


def _teacher(connection: sqlite3.Connection, path: Path, items: list, stamp: str) -> None:
    document = Document()
    _add(document, "교사용 원고", 18, True)
    _add(document, f"문서 번호: {stamp}")
    _add(document, "각 문항은 성취기준 코드와 출처 ID로 데이터베이스와 연결됩니다.")
    for number, item in enumerate(items, start=1):
        standard = require_standard(connection, item["standard_code"])
        standard_source = require_source(connection, standard["source_id"])
        item_source = require_source(connection, item["source_id"])
        _add(document, f"{number}. {item['stem']}", 12, True)
        for choice in json.loads(item["choices_json"]):
            _add(document, choice)
        _add(document, f"정답: {item['answer']}", bold=True)
        _add(document, f"해설: {item['explanation']}")
        _add(document, f"성취기준 {standard['code']} ({standard['school_level']} {standard['subject']} / {standard['unit_title']})")
        _add(document, standard["statement"])
        _add(
            document,
            f"성취기준 출처: {standard_source['issuer']} {standard_source['document_number']} / 출처 ID {standard_source['source_id']}",
        )
        _add(document, f"문항 출처: {item_source['title']} / 출처 ID {item_source['source_id']} ({item['origin']})")
        notes = connection.execute(
            """
            SELECT c.note, c.page_ref, c.note_kind, t.title, t.publisher, t.approval_number
            FROM comparisons c
            JOIN textbooks t ON t.textbook_id = c.textbook_id
            WHERE c.standard_code = ? AND c.note_kind = '교사메모'
            ORDER BY t.title, t.publisher
            """,
            (item["standard_code"],),
        ).fetchall()
        if notes:
            _add(document, "교사 비교 메모", bold=True)
            for note in notes:
                _add(
                    document,
                    f"- {note['title']} / {note['publisher']} / 검정번호 {note['approval_number']} / {note['page_ref']}: {note['note']}",
                )
    document.save(path)
