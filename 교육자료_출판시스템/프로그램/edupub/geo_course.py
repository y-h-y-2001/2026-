"""중학교 지리에서 고등학교 지리 과목으로 이어지는 수업 자료를 만든다."""

from __future__ import annotations

import sqlite3
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from edupub.db import connect
from edupub.geo_maps import build_maps
from edupub.paths import ROOT

OUT = ROOT / "산출물" / "지리연계"


def _run(paragraph, text: str, size: int = 11, bold: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    return run


def _p(document: Document, text: str, size: int = 11, bold: bool = False):
    paragraph = document.add_paragraph()
    _run(paragraph, text, size, bold)
    return paragraph


def _books(connection: sqlite3.Connection, title: str) -> list[str]:
    rows = connection.execute(
        """
        SELECT publisher, lead_author, approval_number
        FROM textbooks WHERE title = ? ORDER BY publisher
        """,
        (title,),
    ).fetchall()
    lines = []
    for row in rows:
        author = row["lead_author"] or "대표 저자 칸 미확인"
        lines.append(f"{row['publisher']} / {author} / 검정번호 {row['approval_number']}")
    return lines


def _cover(document: Document, title: str, subtitle: str):
    _p(document, title, 20, True)
    _p(document, subtitle, 12)
    _p(
        document,
        "이 자료는 아이클라우드의 교과서 파일을 열지 못한 상태에서 만들었습니다. "
        "교과서 문장, 교과서 삽화, 확인하지 않은 통계 수치는 넣지 않았습니다.",
    )
    _p(
        document,
        "성취기준 문장은 교육부 고시 제2022-33호 [별책 7] 사회과 교육과정 한글 파일에서 읽었습니다. "
        "지도는 Natural Earth 퍼블릭 도메인 경계로 새로 그렸고, 위도·경도를 평면으로 펼쳐 면적과 거리가 왜곡됩니다.",
    )


CHAPTERS = [
    {
        "no": 1,
        "title": "위치와 스케일",
        "map": "world",
        "middle": ["9사(지리)01-01", "9사(지리)01-02", "9사(지리)01-03"],
        "high": ["10통사1-01-01", "12세지01-01", "12세지01-03"],
        "body": (
            "지리는 장소를 외우는 과목이 아니라, 위치가 자연환경과 사람 사는 모습과 어떻게 연결되는지 추론하는 과목이다. "
            "같은 일이라도 우리 동네에서 보면 달리고, 우리나라에서 보면 다르고, 세계에서 보면 또 달라진다. 이 범위를 스케일이라고 부른다. "
            "중학교에서는 세계 여러 지역의 특성과, 세계와 지역이 서로 영향을 주고받는 사례를 찾는다. "
            "고등학교 통합사회에서는 같은 현상을 시간·공간·사회·윤리의 관점으로 나누어 보고, "
            "세계시민과 지리에서는 세계화와 지역화를 스케일에 따라 읽는다."
        ),
        "tasks": [
            "지도에서 서울과 런던을 찾고, 어느 쪽이 경도 0도에 더 가까운지 말한다.",
            "우리 학교 주변에서 볼 수 있는 일과, 그것이 세계 다른 지역과 연결될 수 있는 길을 한 가지 적는다. 숫자는 확인한 출처가 있을 때만 적는다.",
        ],
        "questions": [
            {
                "stem": "스케일을 우리 동네, 우리나라, 세계로 바꾸면 같은 현상이 다르게 보일 수 있다. 그 이유를 예시 없이 한 문장으로 쓰시오.",
                "answer": "보는 공간 범위가 달라지면 같은 일의 원인과 영향이 다르게 드러나기 때문이다.",
                "code": "9사(지리)01-02",
            },
            {
                "stem": "세계 한 지역의 특성을 말할 때 위치만 말하면 부족한 이유를, 자연환경과 사람 사는 모습의 관계로 설명하시오.",
                "answer": "위치는 자연환경과 연결되고, 그 환경 속에서 사람 사는 모습이 달라지므로 둘을 함께 보아야 지역의 특성이 된다.",
                "code": "9사(지리)01-01",
            },
        ],
    },
    {
        "no": 2,
        "title": "아시아",
        "map": "asia",
        "middle": ["9사(지리)02-01", "9사(지리)02-02", "9사(지리)02-03", "9사(지리)02-04"],
        "high": ["12세지02-03", "12세지03-01", "12한탐" ],
        "body": (
            "아시아 지도에서는 먼저 우리나라, 동아시아, 동남아시아, 남아시아의 상대적 위치를 잡는다. "
            "서울·도쿄·베이징·상하이는 중위도 동아시아에 모여 있고, 자카르타와 싱가포르는 적도 부근이다. "
            "위도가 다르면 기온과 생활 모습이 달라질 수 있다. 다만 이 지도만으로 기후 수치를 단정하지 않는다. "
            "이 자료의 대륙 구분은 Natural Earth를 따른다. 그 자료는 러시아를 유럽으로 분류하므로, 아시아 지도에서 시베리아는 색칠되지 않는다. "
            "종교 경관, 인구 구조, 산업의 변화는 중학교 아시아 단원의 핵심이고, "
            "고등학교에서는 한국지리 탐구와 세계시민과 지리에서 같은 장소를 더 깊게 조사한다."
        ),
        "tasks": [
            "지도의 도시 여섯 곳을 적도에 가까운 순서대로 가설을 세워 보고, 위도 숫자를 지도의 가로선과 비교해 확인한다.",
            "아시아의 문화경관 사례는 교과서나 확인된 사진 자료가 있을 때만 붙인다. 지금은 칸만 남긴다.",
        ],
        "questions": [
            {
                "stem": "자카르타와 서울 중 적도에 더 가까운 도시를 고르고, 그 위치 차이가 생활 환경에 어떤 차이를 만들 수 있는지 한 가지 추론하시오.",
                "answer": "자카르타가 더 가깝다. 저위도는 연중 기온이 높게 유지되기 쉬워 의복·주거·농사 시기가 중위도 도시와 달라질 수 있다. 구체적인 기온 수치는 기상 자료를 확인한 뒤에만 적는다.",
                "code": "9사(지리)02-01",
            }
        ],
    },
    {
        "no": 3,
        "title": "유럽",
        "map": "europe",
        "middle": ["9사(지리)03-01", "9사(지리)03-02", "9사(지리)03-03"],
        "high": ["12세지01-02", "12도탐03-01"],
        "body": (
            "유럽은 여러 나라가 육지로 붙어 있어 통합과 분리가 도시와 경계에서 함께 보인다. "
            "런던, 파리, 베를린, 모스크바의 위치만으로도 서유럽과 유럽 동부의 거리를 비교할 수 있다. "
            "중학교에서는 유럽 도시의 유형과, 기후위기에 대응하는 도시, 유럽연합처럼 지역이 모이고 나뉘는 움직임을 다룬다. "
            "고등학교 도시의 미래 탐구에서는 환경 문제와 재난을 자연 요인과 사회 요인이 겹친 일로 다시 본다."
        ),
        "tasks": [
            "네 도시 가운데 대서양에 가장 가까운 도시와 가장 안쪽에 있는 도시를 고른다.",
            "유럽의 통합이나 분리 사례를 적으려면 기사나 교과서 쪽수를 함께 적는다. 기억만으로 국가 이름을 단정하지 않는다.",
        ],
        "questions": [
            {
                "stem": "유럽 도시를 지속가능하게 만들려는 노력을 조사할 때, 지도 위에서 먼저 확인해야 할 것 두 가지를 쓰시오.",
                "answer": "그 도시의 위치와, 그 도시가 속한 나라·지역의 자연환경 특성이다.",
                "code": "9사(지리)03-01",
            }
        ],
    },
    {
        "no": 4,
        "title": "아프리카",
        "map": "africa",
        "middle": ["9사(지리)04-01", "9사(지리)04-02", "9사(지리)04-03"],
        "high": ["12세지02-01", "12세지04-02", "12기지02-02"],
        "body": (
            "아프리카는 하나의 모습으로 줄여 말할 수 없다. 카이로, 라고스, 나이로비, 요하네스버그는 대륙의 다른 위도와 해안에 있다. "
            "중학교 성취기준은 자연환경과 문화, 지역의 잠재력을 찾고, 학생이 가지고 있던 아프리카 인식을 되돌아보게 한다. "
            "고등학교에서는 기후를 활용하거나 극복한 사례, 환경 문제, 기후 영향이 지역마다 다르게 나타나는 이유를 이어서 본다."
        ),
        "tasks": [
            "네 도시가 적도, 북회귀선, 남회귀선 중 어디에 가까운지 지도의 선과 비교한다.",
            "‘아프리카는 모두 같다’는 말이 왜 지리에 맞지 않는지, 지도의 위치 차이만으로 두 문장 쓴다.",
        ],
        "questions": [
            {
                "stem": "아프리카에 대한 자신의 기존 인식을 한 가지 적고, 지도에서 확인한 위치 차이가 그 인식을 어떻게 바꾸는지 쓰시오.",
                "answer": "정답은 하나가 아니다. 대륙 전체를 한 기후나 한 생활로 묶는 인식이, 위도와 해안 위치가 다른 도시들을 본 뒤 흔들리면 성취기준에 닿아 있다.",
                "code": "9사(지리)04-02",
            }
        ],
    },
    {
        "no": 5,
        "title": "아메리카",
        "map": "americas",
        "middle": ["9사(지리)05-01", "9사(지리)05-02", "9사(지리)05-03"],
        "high": ["12세지03-03", "12세지02-04"],
        "body": (
            "아메리카는 북아메리카와 남아메리카로 이어진다. 뉴욕, 멕시코시티, 상파울루, 부에노스아이레스는 북쪽에서 남쪽으로 놓여 있다. "
            "중학교에서는 다양한 사람들이 모여 문화가 섞인 모습을 사례로 설명하고, 기업 활동이 어떤 지역에 놓이고 다른 지역으로 옮겨 갈 때 그 지역이 어떻게 바뀌는지 분석한다. "
            "고등학교 세계시민과 지리는 초국적 기업과 글로벌 경제의 공간적 차이를 같은 질문으로 이어 간다. "
            "기업 이름이나 이전 사실은 확인된 자료의 출처를 적기 전에는 문제에 넣지 않는다."
        ),
        "tasks": [
            "네 도시를 위도 순서로 적고, 북반구와 남반구를 나눈다.",
            "문화가 섞인 사례나 기업 입지 사례는 출처 칸이 채워진 뒤에만 발표한다.",
        ],
        "questions": [
            {
                "stem": "아메리카의 문화를 한 나라의 문화로만 설명하면 안 되는 이유를, 성취기준의 표현을 빌려 쓰시오.",
                "answer": "여러 민족과 인종이 함께 살면서 문화가 섞이는 혼종성이 나타나기 때문이다.",
                "code": "9사(지리)05-02",
            }
        ],
    },
    {
        "no": 6,
        "title": "오세아니아와 위도대",
        "map": "oceania",
        "extra_map": "latitude",
        "middle": ["9사(지리)06-01", "9사(지리)06-02", "9사(지리)06-03"],
        "high": ["12세지02-01", "12세지04-01", "12기지01-01", "12기지01-02"],
        "body": (
            "오세아니아 지도의 시드니와 오클랜드는 남반구 중위도에 있다. 태평양의 작은 섬국은 이 축척에서 잘 보이지 않는다. "
            "그래서 두 번째 그림인 위도대 모식도를 함께 본다. 색띠는 관측된 기후 지도가 아니다. 적도와 회귀선, 극권으로 저위도·중위도·고위도를 나눈 틀이다. "
            "극지방의 개발을 둘러싼 이해관계는 중학교에서 의견을 비교하는 활동이고, "
            "고등학교 기후변화 과목에서는 기후를 바라보는 관점이 여러 가지임을 이어서 다룬다."
        ),
        "tasks": [
            "시드니가 남반구인지 북반구인지 적도 선을 기준으로 말한다.",
            "위도대 모식도와 실제 기후가 다를 수 있는 이유를 한 가지 말한다. 해발고도, 해안, 계절풍처럼 위도 이외의 요인을 떠올리면 된다.",
        ],
        "questions": [
            {
                "stem": "위도대 모식도의 색을 그 장소의 실제 기후 값으로 읽으면 안 되는 이유를 쓰시오.",
                "answer": "색띠는 위도만으로 나눈 학습용 틀이고, 관측 기온이나 강수량이 아니기 때문이다.",
                "code": "12세지02-01",
            }
        ],
    },
    {
        "no": 7,
        "title": "우리나라의 위치와 지역",
        "map": "korea",
        "middle": ["9사(지리)07-01", "9사(지리)07-02", "9사(지리)08-01", "9사(지리)08-02", "9사(지리)09-01", "9사(지리)10-01", "9사(지리)11-01"],
        "high": ["12한탐", "12여지02-01"],
        "body": (
            "이 지도의 색은 행정구역 이름이 아니라, 2022 개정 중학교 사회 지리 영역의 단원 묶음이다. "
            "중부는 수도권·강원·충청, 남부는 영남·호남·제주, 북부는 북한의 시도이다. "
            "시도 이름은 Natural Earth 2022년 5월 자료의 한글 이름이라, 이후 바뀐 공식 행정구역 명칭과 다를 수 있다. "
            "독도와 같은 작은 섬은 이 시도 경계 그림에 면으로 나타나지 않을 수 있다. 독도는 대한민국의 영토이며, 위치를 확인할 때는 국토지리정보원 지도를 쓴다. "
            "고등학교 한국지리 탐구는 이 지역 단원을 여행, 산업, 거주, 환경 문제로 다시 깊게 들어간다."
        ),
        "tasks": [
            "자신이 사는 시·도를 지도에서 찾고, 중부·남부·북부 중 어디에 들어가는지 적는다.",
            "우리 지역의 산·강·바다 가운데 가 본 곳을 한 곳 적고, 그 경관이 매력적인 이유를 자신의 경험으로 쓴다.",
        ],
        "questions": [
            {
                "stem": "수도권과 제주도를 이 지도의 색 구분으로 각각 어느 묶음에 놓는지 쓰시오.",
                "answer": "수도권은 중부, 제주도는 남부이다. 이 묶음은 행정구역 개편이 아니라 중학교 지리 단원의 지역 구분이다.",
                "code": "9사(지리)09-01",
            },
            {
                "stem": "북한의 자연환경과 사람 사는 모습을 이야기할 때, 이 지도만으로 알 수 있는 것과 알 수 없는 것을 하나씩 쓰시오.",
                "answer": "시도의 상대적 위치는 알 수 있다. 기온, 인구, 산업 수치와 작은 섬의 정확한 형태는 이 지도만으로 알 수 없다.",
                "code": "9사(지리)11-01",
            },
        ],
    },
    {
        "no": 8,
        "title": "도시, 기후, 그리고 고등학교 선택",
        "map": "scale",
        "middle": ["9사(지리)09-02", "9사(지리)12-01", "9사(지리)12-02", "9사(지리)12-03"],
        "high": ["10통사1-05-01", "12도탐01-01", "12도탐04-01", "12기지04-01", "12여지01-02", "12세지04-02"],
        "body": (
            "도시는 사람이 모여 사는 방식 그 자체이다. 중학교에서는 수도권의 공간 구조와 변화를 보고, "
            "식량·에너지의 소비와 수입처럼 우리 생활이 다른 지역과 연결된 문제를 찾는다. 수입국 이름과 수치는 통계 출처를 확인한 뒤에 적는다. "
            "고등학교에서는 과목이 나뉜다. 통합사회는 생활공간과 사회, 세계시민과 지리는 세계 환경과 평화, "
            "한국지리 탐구는 우리나라 지역, 도시의 미래 탐구는 도시 문제와 공간 정의, "
            "여행지리는 안전한 이동과 장소, 기후변화와 지속가능한 세계는 기후정의와 일상 실천을 맡는다. "
            "중학교 지리의 스케일 그림은 그 과목들이 서로 다른 층위라는 안내도이다."
        ),
        "tasks": [
            "고등학교 과목 다섯 개 가운데 자신이 더 알아보고 싶은 과목을 고르고, 중학교 성취기준 코드 하나를 연결한다.",
            "우리 지역 문제를 지도로 나타내려면 어떤 자료가 더 필요한지 적는다. 자료가 없으면 ‘아직 없음’이라고 적는다.",
        ],
        "questions": [
            {
                "stem": "도시의 미래 탐구와 여행지리가 모두 도시를 다루더라도 질문이 어떻게 다른지, 성취기준의 취지로 한 문장씩 비교하시오.",
                "answer": "도시의 미래 탐구는 도시에서의 삶과 문제, 지속가능성을 묻고, 여행지리는 도시를 여행하며 장소를 새롭게 보고 안전하게 이동하는 방법을 묻는다.",
                "code": "12도탐01-01",
            }
        ],
    },
]


def _standards(connection: sqlite3.Connection, codes: list[str]) -> list[sqlite3.Row]:
    found = []
    for code in codes:
        if code == "12한탐":
            rows = connection.execute(
                "SELECT code, subject, statement FROM standards WHERE subject = '한국지리 탐구' ORDER BY code"
            ).fetchall()
            found.extend(rows[:3])
            continue
        row = connection.execute(
            "SELECT code, subject, statement FROM standards WHERE code = ?",
            (code,),
        ).fetchone()
        if row:
            found.append(row)
    return found


def _student(connection, maps: dict[str, Path]) -> Path:
    document = Document()
    _cover(document, "지리로 잇다", "중학교 사회 지리 영역에서 고등학교 지리 과목까지")
    _p(document, "각 장의 지도 아래에 있는 활동을 노트에 적는다. 정답은 교사용 수업 노트에 있다.", bold=True)
    for chapter in CHAPTERS:
        document.add_page_break()
        _p(document, f"{chapter['no']}. {chapter['title']}", 16, True)
        _p(document, chapter["body"])
        document.add_picture(str(maps[chapter["map"]]), width=Cm(15.5))
        _p(document, "지도 출처: Natural Earth 퍼블릭 도메인. 학습용으로 다시 그렸다.", size=9)
        if chapter.get("extra_map"):
            document.add_picture(str(maps[chapter["extra_map"]]), width=Cm(15.5))
            _p(document, "위도대 그림은 관측 기후가 아닌 모식도이다.", size=9)
        _p(document, "이 장에서 만나는 성취기준", 12, True)
        for row in _standards(connection, chapter["middle"] + chapter["high"]):
            _p(document, f"{row['code']} ({row['subject']}) {row['statement']}")
        _p(document, "지도와 함께 할 일", 12, True)
        for task in chapter["tasks"]:
            _p(document, f"· {task}")
            _p(document, "내 기록: ________________________________")
    path = OUT / "지리_연계_교재.docx"
    document.save(path)
    return path


def _workbook(connection, maps: dict[str, Path]) -> Path:
    document = Document()
    _cover(document, "지리로 잇다 문제집", "지도로 풀고, 성취기준 코드를 남기는 문제")
    number = 1
    for chapter in CHAPTERS:
        document.add_page_break()
        _p(document, f"{chapter['no']}. {chapter['title']}", 16, True)
        document.add_picture(str(maps[chapter["map"]]), width=Cm(14.5))
        for item in chapter["questions"]:
            _p(document, f"{number}. {item['stem']}", 12, True)
            _p(document, f"관련 성취기준: {item['code']}", size=10)
            _p(document, "답: ________________________________________________")
            _p(document, "________________________________________________")
            number += 1
        _p(document, "자료 붙임 칸", 12, True)
        _p(document, "통계나 사례를 넣으면 기관, 자료 이름, 시점을 함께 적는다. 확인 전이면 비운다.")
        _p(document, "출처: ________________  시점: ________________")
    path = OUT / "지리_연계_문제집.docx"
    document.save(path)
    return path


def _teacher(connection, maps: dict[str, Path]) -> Path:
    document = Document()
    _cover(document, "지리로 잇다 수업 노트", "교사용. 차시 흐름, 정답 요지, 검정 도서 서지")
    _p(document, "수업은 한 장에 1차시로 잡을 수 있다. 도입에서 지도를 먼저 읽고, 성취기준 문장을 학생 언어로 줄인 뒤, 빈칸에 기록을 남긴다.", )
    _p(document, "우리 학교 교과서를 연결하는 방법", 14, True)
    _p(document, "아이클라우드나 학교 컴퓨터에 있는 교과서 PDF를 연 다음, 프로그램 메뉴 7번에 검정번호를 적고 파일을 연결한다. 쪽수는 그때 아래 빈칸에 적는다. 공고에 없는 책은 연결되지 않는다.")
    social = _books(connection, "사회 ①")
    _p(document, "중학교 사회 ①에 지리 영역이 들어 있는 검정 도서 서지. 어느 단원이 몇 쪽인지는 파일을 열기 전에는 모른다.", 12, True)
    for line in social:
        _p(document, f"· {line}")
    _p(document, "우리 학교 사회 교과서 쪽: ____________________")
    citizen = _books(connection, "세계시민과 지리")
    _p(document, "고등학교 세계시민과 지리 검정 도서 서지", 12, True)
    for line in citizen:
        _p(document, f"· {line}")
    _p(document, "지리 부도 서지", 12, True)
    for line in _books(connection, "지리 부도"):
        _p(document, f"· {line}")
    _p(
        document,
        "한국지리 탐구, 여행지리, 도시의 미래 탐구, 기후변화와 지속가능한 세계는 2024년 8월 30일 합격 공고에서 사회·역사 도서로 확인되지 않았다. "
        "그 과목의 출판사 이름은 여기에 적지 않는다.",
    )
    for chapter in CHAPTERS:
        document.add_page_break()
        _p(document, f"{chapter['no']}차시. {chapter['title']}", 16, True)
        document.add_picture(str(maps[chapter["map"]]), width=Cm(14))
        _p(document, "흐름", 12, True)
        _p(document, "1. 지도를 화면이나 인쇄물로 보여 주고, 제목과 왜곡 안내를 읽힌다.")
        _p(document, "2. 학생이 도시나 지역을 찾게 한 뒤, 성취기준 문장을 한 줄로 말하게 한다.")
        _p(document, "3. 문제집 답을 받고, 출처가 없는 수치는 지우게 한다.")
        _p(document, "정답 요지", 12, True)
        for item in chapter["questions"]:
            _p(document, f"{item['code']} {item['stem']}", bold=True)
            _p(document, f"요지: {item['answer']}")
        _p(document, "교과서 쪽 메모", 12, True)
        _p(document, "출판사: __________  책 이름: __________  쪽: __________  확인한 날: __________")
    path = OUT / "지리_연계_수업노트.docx"
    document.save(path)
    return path


def build_course() -> list[Path]:
    OUT.mkdir(parents=True, exist_ok=True)
    maps = build_maps()
    connection = connect()
    try:
        paths = [
            _student(connection, maps),
            _workbook(connection, maps),
            _teacher(connection, maps),
        ]
    finally:
        connection.close()
    return paths
